"""Sinh file bao cao Word (.docx) tu noi dung 4 chuong trong thu muc docs/.

Bao cao bam dung khung yeu cau trong CauTrucBaoCao.docx:
    Chuong 1 - Tong quan bai toan
    Chuong 2 - Thiet ke va quan tri CSDL quan he
    Chuong 3 - Luu tru du lieu phi cau truc (NoSQL)
    Chuong 4 - Xu ly va phan tich du lieu lon

Chay:  uv run python -m ecommerce_bigdata.report
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import ASSETS_DIR, DOCS_DIR, PROJECT_ROOT

OUTPUT = PROJECT_ROOT / "BaoCao_ECommerce_BigData.docx"

BRAND = RGBColor(0xEE, 0x4D, 0x2D)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"

# Thong tin sinh vien - sua truc tiep o day truoc khi nop
STUDENT = {
    "truong": "TRUONG DAI HOC ...",
    "khoa": "KHOA CONG NGHE THONG TIN",
    "mon": "CO SO DU LIEU & BIG DATA",
    "de_tai": "KIẾN TRÚC DỮ LIỆU KÉP CHO HỆ THỐNG THƯƠNG MẠI ĐIỆN TỬ",
    "phu_de": "Sàn giao dịch đa người bán — SQL Server và MongoDB",
    "sinh_vien": "......................................",
    "ma_so": "......................................",
    "giang_vien": "......................................",
    "nam_hoc": "2025 – 2026",
}

CHAPTERS = [
    ("01_yeu_cau_nghiep_vu.md", "CHƯƠNG 1. TỔNG QUAN BÀI TOÁN"),
    ("02_thiet_ke_csdl.md", "CHƯƠNG 2. THIẾT KẾ VÀ QUẢN TRỊ CSDL QUAN HỆ"),
    ("03_nosql.md", "CHƯƠNG 3. LƯU TRỮ DỮ LIỆU PHI CẤU TRÚC (NoSQL)"),
    ("04a_hadoop.md", "CHƯƠNG 4. XỬ LÝ VÀ PHÂN TÍCH DỮ LIỆU LỚN"),
    ("04b_spark.md", None),          # None = noi tiep chuong trên, khong sang chuong moi
    ("04_phan_tich.md", None),
]


# ---------------------------------------------------------------------------
# Bo cong cu dinh dang
# ---------------------------------------------------------------------------
def set_cell_background(cell, hex_color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def add_field(paragraph, instruction: str) -> None:
    """Chen truong dong (vi du so trang, muc luc) vao doan van."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, end):
        run._r.append(node)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for level, size, color in ((1, 16, BRAND), (2, 14, INK), (3, 13, INK)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)


# ---------------------------------------------------------------------------
# Phan tich Markdown
# ---------------------------------------------------------------------------
@dataclass
class Block:
    kind: str                  # heading | para | bullet | numbered | table | image | code
    content: object
    level: int = 0


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    lines = text.splitlines()
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # --- Khoi ma nguon
        if stripped.startswith("```"):
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", "\n".join(code)))
            continue

        # --- Tieu de
        if m := re.match(r"^(#{1,4})\s+(.*)$", stripped):
            blocks.append(Block("heading", m.group(2).strip(), len(m.group(1))))
            i += 1
            continue

        # --- Anh
        if m := re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped):
            blocks.append(Block("image", (m.group(1), m.group(2))))
            i += 1
            continue

        # --- Bang
        if stripped.startswith("|") and i + 1 < len(lines) and \
           re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = lines[i].strip()
                if not re.match(r"^\|[\s:\-|]+\|$", row):
                    rows.append([c.strip() for c in row.strip("|").split("|")])
                i += 1
            blocks.append(Block("table", rows))
            continue

        # --- Danh sach
        if m := re.match(r"^[-*]\s+(.*)$", stripped):
            blocks.append(Block("bullet", m.group(1)))
            i += 1
            continue
        if m := re.match(r"^(\d+)\.\s+(.*)$", stripped):
            blocks.append(Block("numbered", m.group(2)))
            i += 1
            continue

        # --- Doan van (gop cac dong lien tiep)
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,4}\s|[-*]\s|\d+\.\s|\||!\[|```)", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        blocks.append(Block("para", " ".join(para)))

    return blocks


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_rich_text(paragraph, text: str) -> None:
    """Xu ly **dam**, *nghieng*, `ma nguon` trong mot doan."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(11)
            run.font.color.rgb = BRAND
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


# ---------------------------------------------------------------------------
# Dung tai lieu
# ---------------------------------------------------------------------------
def render_block(doc: Document, block: Block, base_dir: Path) -> None:
    if block.kind == "heading":
        # Tieu de cap 1 trong file md la ten chuong, da in rieng -> bo qua.
        # '##' trong Markdown phai thanh Heading 2 de nam duoi ten chuong
        # (Heading 1), neu khong muc luc se phang.
        if block.level == 1:
            return
        doc.add_heading(block.content, min(block.level, 4))

    elif block.kind == "para":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich_text(p, block.content)

    elif block.kind in ("bullet", "numbered"):
        style = "List Bullet" if block.kind == "bullet" else "List Number"
        p = doc.add_paragraph(style=style)
        add_rich_text(p, block.content)

    elif block.kind == "code":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run(block.content)
        run.font.name = CODE_FONT
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    elif block.kind == "table":
        rows = block.content
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                if c >= len(table.columns):
                    continue
                cell = table.cell(r, c)
                cell.text = ""
                p = cell.paragraphs[0]
                add_rich_text(p, text)
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    if r == 0:
                        run.font.bold = True
                if r == 0:
                    set_cell_background(cell, "FFE8E1")
        doc.add_paragraph()

    elif block.kind == "image":
        alt, rel_path = block.content
        img = (base_dir / rel_path).resolve()
        if not img.exists():
            return
        doc.add_picture(str(img), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = doc.add_paragraph(f"Hình: {alt}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.font.size = Pt(10.5)
            run.italic = True
            run.font.color.rgb = MUTED


def build_cover(doc: Document) -> None:
    def centered(text: str, size: int, bold: bool = False,
                 color: RGBColor = INK, space: int = 6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return p

    centered(STUDENT["truong"], 13, True)
    centered(STUDENT["khoa"], 13, True, space=90)
    centered("BÁO CÁO BÀI TẬP LỚN", 22, True, BRAND)
    centered(f"Môn học: {STUDENT['mon']}", 13, space=40)
    centered("ĐỀ TÀI", 12, True, MUTED)
    centered(STUDENT["de_tai"], 17, True, BRAND)
    centered(STUDENT["phu_de"], 12.5, color=MUTED, space=100)

    for label, key in (("Sinh viên thực hiện", "sinh_vien"), ("Mã số sinh viên", "ma_so"),
                       ("Giảng viên hướng dẫn", "giang_vien"), ("Năm học", "nam_hoc")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{label}: ")
        run.font.size = Pt(13)
        run2 = p.add_run(STUDENT[key])
        run2.font.size = Pt(13)
        run2.font.bold = True

    centered(f"Tháng {date.today().month} năm {date.today().year}", 12,
             color=MUTED, space=0)
    doc.paragraphs[-1].paragraph_format.space_before = Pt(60)
    doc.add_page_break()


def build_toc(doc: Document) -> None:
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("MỤC LỤC")
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = BRAND

    add_field(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u')

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("(Mở file trong Word, bấm Ctrl+A rồi F9 để cập nhật mục lục)")
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = MUTED
    doc.add_page_break()


def add_page_numbers(doc: Document) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")
    for run in footer.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = MUTED


def build() -> Path:
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)

    build_cover(doc)
    build_toc(doc)

    for filename, title in CHAPTERS:
        path = DOCS_DIR / filename
        if not path.exists():
            print(f"   [bo qua] thieu {filename}")
            continue
        if title is not None:
            doc.add_heading(title, 1)
        for block in parse_markdown(path.read_text(encoding="utf-8")):
            render_block(doc, block, path.parent)
        doc.add_page_break()
        print(f"   [{'chuong' if title else 'phan '}] {title or path.stem}")

    # --- Chuong ket luan
    doc.add_heading("KẾT LUẬN", 1)
    for para in (
        "Đồ án đã xây dựng hoàn chỉnh một kiến trúc dữ liệu kép cho hệ thống thương "
        "mại điện tử: dữ liệu giao dịch nằm trong CSDL quan hệ với ràng buộc toàn vẹn "
        "chặt chẽ, dữ liệu hành vi tần suất cao nằm trong MongoDB với schema linh hoạt.",
        "Kết quả đạt được gồm: lược đồ quan hệ 12 bảng chuẩn hóa tới 3NF kèm 4 trigger "
        "nghiệp vụ; 10 view/procedure/function đáp ứng đúng 10 yêu cầu truy vấn đặt ra "
        "ở Chương 1; hai collection MongoDB với khoảng 50.000 sự kiện clickstream và "
        "22.000 log quét mã vạch; và script Python gộp hai nguồn để dựng 6 biểu đồ "
        "phân tích.",
        "Giá trị lớn nhất của kiến trúc kép thể hiện ở Chương 4: những phát hiện như "
        "nhóm sản phẩm được xem nhiều nhưng bán kém, hay quan hệ giữa số chặng trung "
        "chuyển và thời gian giao hàng, chỉ có được khi đối chiếu đồng thời dữ liệu "
        "giao dịch và dữ liệu hành vi — không nguồn nào một mình trả lời được.",
        "Hướng phát triển tiếp theo: đưa luồng sự kiện qua Kafka thay vì ghi trực tiếp "
        "vào MongoDB để chịu tải tốt hơn, và xây dựng kho dữ liệu riêng cho báo cáo "
        "thay vì truy vấn thẳng trên CSDL vận hành.",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.add_run(para)

    add_page_numbers(doc)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(">> Dung bao cao ...")
    path = build()
    size_kb = path.stat().st_size / 1024
    print(f"\n>> Da tao {path.name} ({size_kb:,.0f} KB)")
